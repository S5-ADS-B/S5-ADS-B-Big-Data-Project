#!/usr/bin/env python3
"""Convert the report CSV produced by build_report.sh into a formatted .xlsx file.

Usage: csv_to_xlsx.py input.csv output.xlsx
"""
import csv
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = {
    "group": "Group",
    "roll_no": "Roll No",
    "name": "Name",
    "topic": "Topic",
    "link": "Link",
}


def main(src: str, dst: str) -> None:
    with open(src, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    header = [HEADERS.get(h, h) for h in rows[0]]
    ws.append(header)

    for row in rows[1:]:
        group, roll, name, topic, link = row
        ws.append([int(group) if group.isdigit() else group, roll, name, topic, link])

    # Keep roll numbers as text so values like 007 are not turned into 7
    for cell in ws["B"][1:]:
        cell.number_format = "@"
        cell.data_type = "s"

    # Header styling
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(horizontal="left", vertical="center")

    for cell in ws["A"][1:]:
        cell.alignment = Alignment(horizontal="left")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Column widths based on content
    for idx, col in enumerate(ws.columns, start=1):
        longest = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(idx)].width = min(max(longest + 2, 8), 70)

    wb.save(dst)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("Usage: csv_to_xlsx.py input.csv output.xlsx")
    main(sys.argv[1], sys.argv[2])